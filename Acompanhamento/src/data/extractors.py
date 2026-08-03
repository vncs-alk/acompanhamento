import pandas as pd
import numpy as np
import warnings

class ReadData():
    def __init__(self, caminho_csv, cnpjs_master, cnpjs_gd, cnpjs_fics):
        self.caminho_csv = caminho_csv
        self.cnpjs_master = cnpjs_master
        self.cnpjs_gd = cnpjs_gd
        self.cnpjs_fics = cnpjs_fics
        
        self.cnpjs_todos = self.cnpjs_master + self.cnpjs_gd + self.cnpjs_fics
        self.carregar_e_filtrar()

    def carregar_e_filtrar(self):
        df = pd.read_csv(self.caminho_csv, index_col=0)
        df['DT_COMPTC'] = pd.to_datetime(df['DT_COMPTC'])
        
        if self.cnpjs_todos:
            df = df[df['CNPJ_FUNDO_CLASSE'].isin(self.cnpjs_todos)]
            
        df['CAPTC_LIQ'] = (df['CAPTC_DIA'] - df['RESG_DIA']) / 1000000
        df['PL_MM'] = df['VL_PATRIM_LIQ'] / 1000000
        
        self.dados = df


class ReadExcel():
    def __init__(self, caminho_excel="base_acompanhamento.xlsm", aba="Optimus"):
        import openpyxl
        import shutil
        import tempfile
        import pandas as pd
        
        with tempfile.NamedTemporaryFile(suffix='.xlsm', delete=False) as tmp:
            shutil.copyfile(caminho_excel, tmp.name)
            caminho_leitura = tmp.name
        
        self.dados = pd.read_excel(caminho_leitura, sheet_name=aba, header=None)
        self.nome_estrategia = aba
        
        masters = self.dados[self.dados[0] == "Masters"].index[0]
        gd = self.dados[self.dados[0] == "GD"].index[0]
        fics = self.dados[self.dados[0] == "Fics"].index[0]

        self.infos_masters = self.dados.iloc[masters + 1:gd].dropna(how='all').dropna(axis=1, how='all')
        if len(self.infos_masters) > 0: self.infos_masters.columns = self.dados.iloc[masters].dropna()

        self.infos_gd = self.dados.iloc[gd + 1:fics].dropna(how='all').dropna(axis=1, how='all')
        if len(self.infos_gd) > 0: self.infos_gd.columns = self.dados.iloc[gd].dropna()

        self.infos_fics = self.dados.iloc[fics + 1:].dropna(how='all').dropna(axis=1, how='all')
        if len(self.infos_fics) > 0: self.infos_fics.columns = self.dados.iloc[fics].dropna()

        self.cnpjs_prateleira_azul = []
        wb = openpyxl.load_workbook(caminho_leitura, read_only=True, data_only=True)
        if 'Aba_Auxiliar' in wb.sheetnames:
            ws_aux = wb['Aba_Auxiliar']
            for row in ws_aux.iter_rows(min_col=1, max_col=1, values_only=True):
                if row[0]:
                    self.cnpjs_prateleira_azul.append(str(row[0]).strip())
        wb.close()

    def extrair_listas(self):
        import pandas as pd
        dicionario_geral = {}
        dict_vol_target = {}

        def processar_df(df_fundo):
            lista = []
            dic = {}
            if df_fundo is None or df_fundo.empty: return lista, dic
            
            for idx, row in df_fundo.iterrows():
                if pd.isna(row.iloc[0]): continue
                cnpj = str(row.iloc[0]).strip()
                nome = str(row.iloc[1]).strip() if len(row) > 1 and pd.notnull(row.iloc[1]) else ""
                vol_target = float(row.iloc[2]) if len(row) > 2 and pd.notnull(row.iloc[2]) else 0.0
                taxa_adm = float(row.iloc[3]) if len(row) > 3 and pd.notnull(row.iloc[3]) else 0.02
                taxa_perf = float(row.iloc[4]) if len(row) > 4 and pd.notnull(row.iloc[4]) else 0.2
                
                lista.append((cnpj, taxa_adm, taxa_perf))
                dic[cnpj] = nome
                dict_vol_target[cnpj] = vol_target
                dicionario_geral[cnpj] = nome
            return lista, dic

        lista_masters, dicionario_masters = processar_df(self.infos_masters)
        lista_gd, _ = processar_df(self.infos_gd)
        lista_fics, dicionario_fics = processar_df(self.infos_fics)
        
        if self.cnpjs_prateleira_azul:
            dicionario_prateleira = {cnpj: dicionario_geral.get(cnpj, cnpj) for cnpj in self.cnpjs_prateleira_azul}
        else:
            dicionario_prateleira = dicionario_geral
            
        return lista_masters, lista_gd, lista_fics, dicionario_prateleira, dicionario_masters, dicionario_fics, dict_vol_target

